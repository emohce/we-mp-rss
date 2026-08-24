from __future__ import annotations

import html
import io
import json
import re
from dataclasses import dataclass
from datetime import datetime

from bs4 import BeautifulSoup


@dataclass(frozen=True)
class ExportArtifact:
    filename: str
    media_type: str
    content: bytes


def safe_filename(value: str, fallback: str = "article") -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|\x00-\x1f]", "_", value or "")
    cleaned = re.sub(r"\s+", " ", cleaned).strip(" ._")
    return (cleaned or fallback)[:100]


def _text_content(article: dict) -> str:
    source = str(article.get("content") or article.get("content_html") or "")
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", source)).strip()


def _safe_web_url(value: object) -> str:
    url = str(value or "").strip()
    return url if re.match(r"^https?://", url, flags=re.I) else "#"


def _safe_html_content(article: dict) -> str:
    source = str(article.get("content_html") or article.get("content") or "")
    soup = BeautifulSoup(source, "html.parser")
    for tag in soup.find_all(
        ["script", "iframe", "object", "embed", "base", "meta", "form", "link"]
    ):
        tag.decompose()
    for tag in soup.find_all(True):
        for attribute in list(tag.attrs):
            normalized = attribute.lower()
            value = tag.attrs.get(attribute)
            if normalized.startswith("on") or normalized in {"srcdoc", "formaction"}:
                del tag.attrs[attribute]
            elif normalized == "href" and not re.match(r"^(https?://|#)", str(value), flags=re.I):
                tag.attrs[attribute] = "#"
            elif normalized in {"src", "poster"} and not re.match(
                r"^(https?://|data:image/)", str(value), flags=re.I
            ):
                del tag.attrs[attribute]
    return str(soup)


class SingleArticleExporter:
    FORMATS = {"md", "html", "json", "pdf", "docx"}

    def export(self, article: dict, format_name: str) -> ExportArtifact:
        normalized = format_name.lower().lstrip(".")
        if normalized not in self.FORMATS:
            raise ValueError(f"unsupported single article format: {normalized}")
        title = str(article.get("title") or "Untitled")
        base = safe_filename(title, f"article-{article.get('id', '')}")
        if normalized == "md":
            body = self._markdown(article).encode("utf-8")
            return ExportArtifact(f"{base}.md", "text/markdown; charset=utf-8", body)
        if normalized == "html":
            body = self._html(article).encode("utf-8")
            return ExportArtifact(f"{base}.html", "text/html; charset=utf-8", body)
        if normalized == "json":
            body = json.dumps(article, ensure_ascii=False, indent=2, default=str).encode("utf-8")
            return ExportArtifact(f"{base}.json", "application/json", body)
        if normalized == "pdf":
            return ExportArtifact(f"{base}.pdf", "application/pdf", self._pdf(article))
        return ExportArtifact(
            f"{base}.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            self._docx(article),
        )

    @staticmethod
    def _markdown(article: dict) -> str:
        metadata = {
            "id": article.get("id"),
            "title": article.get("title"),
            "author": article.get("author", ""),
            "source": article.get("mp_name") or article.get("mp_id", ""),
            "publish_time": article.get("publish_time"),
            "source_url": article.get("url") or article.get("link", ""),
        }
        frontmatter = "\n".join(
            f"{key}: {json.dumps(value, ensure_ascii=False, default=str)}" for key, value in metadata.items()
        )
        content = _safe_html_content(article)
        return f"---\n{frontmatter}\n---\n\n# {article.get('title') or 'Untitled'}\n\n{content}\n"

    @staticmethod
    def _html(article: dict) -> str:
        title = html.escape(str(article.get("title") or "Untitled"))
        source_url = html.escape(
            _safe_web_url(article.get("url") or article.get("link")), quote=True
        )
        body = _safe_html_content(article)
        return (
            "<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'>"
            "<meta http-equiv='Content-Security-Policy' content=\"default-src 'none'; "
            "img-src https: data:; media-src https:; style-src 'unsafe-inline'; "
            "base-uri 'none'; form-action 'none'\">"
            f"<title>{title}</title><meta name='viewport' content='width=device-width,initial-scale=1'>"
            "<style>body{max-width:760px;margin:40px auto;padding:0 20px;font:16px/1.75 system-ui;color:#202124}"
            "img{max-width:100%;height:auto}.meta{color:#68707b;font-size:14px}</style></head><body>"
            f"<h1>{title}</h1><p class='meta'><a href='{source_url}' rel='noopener noreferrer'>查看原文</a></p><article>{body}</article>"
            "</body></html>"
        )

    @staticmethod
    def _pdf(article: dict) -> bytes:
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
        except ImportError as exc:
            raise RuntimeError("reportlab is required for PDF export") from exc
        buffer = io.BytesIO()
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        styles = getSampleStyleSheet()
        styles["Title"].fontName = "STSong-Light"
        styles["BodyText"].fontName = "STSong-Light"
        document = SimpleDocTemplate(buffer, pagesize=A4)
        story = [Paragraph(html.escape(str(article.get("title") or "Untitled")), styles["Title"]), Spacer(1, 12)]
        text = html.escape(_text_content(article)).replace("\n", "<br/>")
        for offset in range(0, len(text), 2000):
            story.append(Paragraph(text[offset : offset + 2000], styles["BodyText"]))
            story.append(Spacer(1, 8))
        document.build(story)
        return buffer.getvalue()

    @staticmethod
    def _docx(article: dict) -> bytes:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("python-docx is required for DOCX export") from exc
        document = Document()
        document.add_heading(str(article.get("title") or "Untitled"), level=0)
        source = str(article.get("mp_name") or article.get("mp_id") or "")
        if source:
            document.add_paragraph(source)
        text = _text_content(article)
        for offset in range(0, len(text), 4000):
            document.add_paragraph(text[offset : offset + 4000])
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()
